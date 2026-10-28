"""
Seed script to initialize data/jobs.xlsx with the required sheets and realistic sample data.
"""
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DEFAULT_EXCEL_PATH = Path("data/jobs.xlsx")

def create_initial_workbook(target_path: Path = DEFAULT_EXCEL_PATH):
    target_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    # remove default sheet
    wb.remove(wb.active)

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    # 1. Jobs Sheet
    ws_jobs = wb.create_sheet(title="Jobs")
    jobs_headers = [
        "JobID", "EmployerID", "JobType", "Location", "Locality", "District",
        "State", "Salary", "SalaryType", "Hours", "RequiredExperience",
        "Vacancies", "Description", "Status", "Verified", "CreatedAt", "ExpiryDate"
    ]
    ws_jobs.append(jobs_headers)

    jobs_data = [
        [
            "JOB-1001", "EMP-101", "Driver", "Hyderabad", "Uppal", "Medchal-Malkajgiri",
            "Telangana", 18000, "Month", "9 AM - 6 PM", "1 year",
            2, "Commercial car driver needed for local trips around Uppal and Secunderabad.", "Active", "Verified", "2026-10-01", "2026-12-31"
        ],
        [
            "JOB-1002", "EMP-102", "Construction", "Hyderabad", "Dilsukhnagar", "Hyderabad",
            "Telangana", 750, "Day", "8 AM - 5 PM", "Fresher",
            5, "Masonry and site helpers required for residential apartment project in Dilsukhnagar.", "Active", "Verified", "2026-10-05", "2026-12-15"
        ],
        [
            "JOB-1003", "EMP-103", "Electrician", "Hyderabad", "Secunderabad", "Hyderabad",
            "Telangana", 20000, "Month", "10 AM - 7 PM", "2 years",
            1, "Commercial wiring, panel maintenance, and appliance repair electrician.", "Active", "Verified", "2026-10-10", "2026-12-20"
        ],
        [
            "JOB-1004", "EMP-104", "Cleaner", "Hyderabad", "Kukatpally", "Medchal-Malkajgiri",
            "Telangana", 14000, "Month", "7 AM - 3 PM", "Fresher",
            4, "Deep cleaning and housekeeping staff needed for offices and retail stores.", "Active", "Verified", "2026-10-12", "2026-12-31"
        ],
        [
            "JOB-1005", "EMP-105", "Shop / Helper", "Hyderabad", "Uppal", "Medchal-Malkajgiri",
            "Telangana", 15000, "Month", "8 AM - 8 PM", "Fresher",
            2, "Supermarket helper for counter sales, stock arrangement, and customer delivery.", "Active", "Verified", "2026-10-15", "2026-12-31"
        ],
        [
            "JOB-1006", "EMP-101", "Driver", "Hyderabad", "Secunderabad", "Hyderabad",
            "Telangana", 19000, "Month", "8 AM - 5 PM", "2 years",
            1, "Delivery van driver for wholesale logistics depot in Secunderabad.", "Active", "Verified", "2026-10-18", "2026-12-31"
        ],
        [
            "JOB-1007", "EMP-106", "Driver", "Hyderabad", "Uppal", "Medchal-Malkajgiri",
            "Telangana", 25000, "Month", "Flexible", "5 years",
            1, "High pay driver role pending administrative verification.", "Active", "Pending", "2026-10-25", "2026-12-31"
        ],
        [
            "JOB-1008", "EMP-101", "Driver", "Hyderabad", "Uppal", "Medchal-Malkajgiri",
            "Telangana", 16000, "Month", "9 AM - 6 PM", "1 year",
            1, "Position closed as vacancy filled.", "Closed", "Verified", "2026-09-01", "2026-12-31"
        ],
        [
            "JOB-1009", "EMP-102", "Construction", "Hyderabad", "Uppal", "Medchal-Malkajgiri",
            "Telangana", 700, "Day", "8 AM - 5 PM", "Fresher",
            2, "Old foundation work completed and expired.", "Active", "Verified", "2026-08-01", "2026-09-01"
        ]
    ]
    for row in jobs_data:
        ws_jobs.append(row)

    # 2. Employers Sheet
    ws_employers = wb.create_sheet(title="Employers")
    emp_headers = ["EmployerID", "CompanyName", "ContactName", "Phone", "Location", "Verified", "CreatedAt"]
    ws_employers.append(emp_headers)
    emp_data = [
        ["EMP-101", "Sri Krishna Logistics", "Ramesh Reddy", "+919876543210", "Uppal, Hyderabad", "Verified", "2026-09-15"],
        ["EMP-102", "Balaji Constructions", "Srinivas Rao", "+919876543211", "Dilsukhnagar, Hyderabad", "Verified", "2026-09-20"],
        ["EMP-103", "Metro Electricals", "Md. Akbar Ali", "+919876543212", "Secunderabad, Hyderabad", "Verified", "2026-09-22"],
        ["EMP-104", "Swachh Facility Services", "Anitha Devi", "+919876543213", "Kukatpally, Hyderabad", "Verified", "2026-09-25"],
        ["EMP-105", "Sri Ram Supermarket", "Venkat Narayana", "+919876543214", "Uppal, Hyderabad", "Verified", "2026-10-01"],
        ["EMP-106", "Quick Shifts Unverified", "Raj Kumar", "+919876543215", "Uppal, Hyderabad", "Pending", "2026-10-24"]
    ]
    for row in emp_data:
        ws_employers.append(row)

    # 3. Workers Sheet
    ws_workers = wb.create_sheet(title="Workers")
    worker_headers = ["WorkerID", "Phone", "Language", "Location", "Locality", "PreferredJobType", "Experience", "CreatedAt"]
    ws_workers.append(worker_headers)
    worker_data = [
        ["WRK-501", "+919848012345", "te", "Hyderabad", "Uppal", "Driver", "3 years", "2026-10-02"],
        ["WRK-502", "+919848054321", "hi", "Hyderabad", "Secunderabad", "Electrician", "2 years", "2026-10-03"],
        ["WRK-503", "+919848098765", "en", "Hyderabad", "Kukatpally", "Cleaner", "Fresher", "2026-10-04"]
    ]
    for row in worker_data:
        ws_workers.append(row)

    # 4. Applications Sheet
    ws_apps = wb.create_sheet(title="Applications")
    app_headers = ["ApplicationID", "WorkerID", "WorkerPhone", "JobID", "AppliedAt", "Status"]
    ws_apps.append(app_headers)
    app_data = [
        ["APP-10001", "WRK-501", "+919848012345", "JOB-1006", "2026-10-20 11:15:00", "Interested"]
    ]
    for row in app_data:
        ws_apps.append(row)

    # 5. Categories Sheet
    ws_cats = wb.create_sheet(title="Categories")
    cat_headers = ["CategoryID", "Name", "TeluguName", "HindiName", "EnglishName", "Active"]
    ws_cats.append(cat_headers)
    cat_data = [
        ["CAT-1", "Driver", "డ్రైవర్", "ड्राइवर", "Driver", True],
        ["CAT-2", "Construction", "భవన నిర్మాణం", "निर्माण कार्य", "Construction", True],
        ["CAT-3", "Electrician", "ఎలక్ట్రీషియన్", "इलेक्ट्रीशियन", "Electrician", True],
        ["CAT-4", "Cleaner", "క్లీనర్", "सफाई कर्मचारी", "Cleaner", True],
        ["CAT-5", "Shop / Helper", "షాప్ సహాయకుడు", "दुकान सहायक", "Shop / Helper", True],
        ["CAT-6", "Plumber", "ప్లంబర్", "प्लंबर", "Plumber", True],
        ["CAT-7", "Security Guard", "సెక్యూరిటీ గార్డ్", "सुरक्षा गार्ड", "Security Guard", True],
        ["CAT-8", "Delivery", "డెలివరీ", "डिलीवरी", "Delivery", True],
        ["CAT-9", "Welder", "వెల్డర్", "वेल्डर", "Welder", True],
        ["CAT-10", "Carpenter", "కార్పెంటర్", "बढ़ई", "Carpenter", True],
        ["CAT-11", "Painter", "పెయింటర్", "पेंटर", "Painter", True],
        ["CAT-12", "Factory Worker", "ఫ్యాక్టరీ కార్మికుడు", "फैक्ट्री वर्कर", "Factory Worker", True],
        ["CAT-13", "Farm Worker", "వ్యవసాయ కార్మికుడు", "खेतिहर मजदूर", "Farm Worker", True],
        ["CAT-14", "Hotel Worker", "హోటల్ వర్కర్", "होटल कर्मचारी", "Hotel Worker", True],
        ["CAT-15", "Mechanic", "మెకానిక్", "मैकेनिक", "Mechanic", True],
        ["CAT-16", "Tailor", "టైలర్", "दर्जी", "Tailor", True],
        ["CAT-17", "Cook", "వంట మనిషి", "रसोइया", "Cook", True],
        ["CAT-18", "Other", "ఇతర పనులు", "अन्य काम", "Other", True]
    ]
    for row in cat_data:
        ws_cats.append(row)

    # Style all sheets
    for ws in wb.worksheets:
        for col_idx, col in enumerate(ws.iter_cols(min_row=1, max_row=1), start=1):
            cell = col[0]
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align

        for row in ws.iter_rows(min_row=1):
            for cell in row:
                cell.border = thin_border
                if cell.row > 1 and not cell.alignment.horizontal:
                    cell.alignment = left_align

        # auto-fit column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(target_path)
    return target_path

if __name__ == "__main__":
    p = create_initial_workbook()
    print(f"Created Excel database at {p.resolve()}")
